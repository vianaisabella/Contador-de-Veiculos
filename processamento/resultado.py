"""
Resultado de uma contagem: salvar/abrir (JSON + recortes) e exportar Excel.
"""

import json
import os
from datetime import datetime

from .config import (
    ConfigProcessamento, formatar_tempo, nome_exibicao, ordenar_classes, pasta_resultados,
)


class Resultado:

    def __init__(self, video, fps, total_frames, poligono, linha, classes,
                 config, modelo, eventos=None, pasta=None, duracao_processamento=None):
        self.video = video
        self.fps = float(fps)
        self.total_frames = int(total_frames)
        self.poligono = [list(map(int, p)) for p in poligono]
        self.linha = [list(map(float, p)) for p in linha]
        self.classes = list(classes)
        self.config = config
        self.modelo = modelo
        self.eventos = eventos or []
        self.pasta = pasta
        self.duracao_processamento = duracao_processamento

    # -------------------------------------------------------------
    # PROPRIEDADES
    # -------------------------------------------------------------

    @property
    def duracao_video(self):
        return self.total_frames / self.fps if self.fps else 0.0

    @property
    def eventos_validos(self):
        return [e for e in self.eventos if not e.get("descartado")]

    def duvidas(self, so_pendentes=False):
        lista = [e for e in self.eventos if e.get("duvida")]
        if so_pendentes:
            lista = [e for e in lista if not e.get("revisado")]
        return lista

    def classes_relatorio(self):
        """Classes do modelo + as que aparecem nos eventos (ex.: não classificado)."""
        return ordenar_classes(set(self.classes) | {e["classe"] for e in self.eventos_validos})

    def totais(self):
        tot = {c: 0 for c in self.classes_relatorio()}
        for e in self.eventos_validos:
            tot[e["classe"]] = tot.get(e["classe"], 0) + 1
        return tot

    def caminho_recorte(self, evento):
        if not evento.get("recorte") or not self.pasta:
            return None
        caminho = os.path.join(self.pasta, evento["recorte"])
        return caminho if os.path.isfile(caminho) else None

    def adicionar_manual(self, tempo, classe):
        proximo = max([e["id"] for e in self.eventos] + [0]) + 1
        if classe not in self.classes:
            self.classes = ordenar_classes(self.classes + [classe])
        evento = {
            "id": proximo, "trilha": None,
            "frame": int(round(tempo * self.fps)), "tempo": round(float(tempo), 2),
            "classe": classe, "classe_sugerida": None, "votos": {},
            "confianca": None, "deteccoes": 0, "sentido": 0, "caixa": None,
            "duvida": False, "motivos": ["adicionado manualmente"],
            "origem": "manual", "revisado": True, "descartado": False, "recorte": None,
        }
        self.eventos.append(evento)
        self.eventos.sort(key=lambda e: e["tempo"])
        return evento

    # -------------------------------------------------------------
    # RESUMO POR INTERVALO
    # -------------------------------------------------------------

    def resumo_intervalos(self, intervalo=None):
        intervalo = intervalo or self.config.intervalo_resumo
        n = max(1, int(-(-self.duracao_video // intervalo)))  # teto
        classes = self.classes_relatorio()
        linhas = []
        for i in range(n):
            inicio = i * intervalo
            fim = min((i + 1) * intervalo, self.duracao_video) if i == n - 1 else (i + 1) * intervalo
            linhas.append({
                "inicio": inicio, "fim": fim,
                "contagem": {c: 0 for c in classes},
            })
        for e in self.eventos_validos:
            i = min(n - 1, int(e["tempo"] // intervalo))
            linhas[i]["contagem"][e["classe"]] = linhas[i]["contagem"].get(e["classe"], 0) + 1
        return linhas

    # -------------------------------------------------------------
    # SALVAR / ABRIR
    # -------------------------------------------------------------

    @staticmethod
    def criar_pasta(video):
        nome = os.path.splitext(os.path.basename(video))[0]
        carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
        pasta = os.path.join(pasta_resultados(), f"{nome}_{carimbo}")
        os.makedirs(os.path.join(pasta, "recortes"), exist_ok=True)
        return pasta

    def salvar_recortes(self, recortes):
        """recortes: dict id_evento -> bytes JPG"""
        pasta_rec = os.path.join(self.pasta, "recortes")
        os.makedirs(pasta_rec, exist_ok=True)
        for e in self.eventos:
            jpg = recortes.get(e["id"])
            if jpg:
                nome = f"evento_{e['id']:05d}.jpg"
                with open(os.path.join(pasta_rec, nome), "wb") as f:
                    f.write(jpg)
                e["recorte"] = f"recortes/{nome}"

    def salvar(self):
        dados = {
            "versao": 1,
            "video": self.video,
            "fps": self.fps,
            "total_frames": self.total_frames,
            "poligono": self.poligono,
            "linha": self.linha,
            "classes": self.classes,
            "config": self.config.to_dict(),
            "modelo": self.modelo,
            "duracao_processamento": self.duracao_processamento,
            "salvo_em": datetime.now().isoformat(timespec="seconds"),
            "eventos": self.eventos,
        }
        caminho = os.path.join(self.pasta, "resultado.json")
        temporario = caminho + ".tmp"
        with open(temporario, "w", encoding="utf-8") as f:
            json.dump(dados, f, ensure_ascii=False, indent=1)
        os.replace(temporario, caminho)
        return caminho

    @classmethod
    def abrir(cls, caminho_json):
        with open(caminho_json, encoding="utf-8") as f:
            dados = json.load(f)
        res = cls(
            video=dados["video"], fps=dados["fps"], total_frames=dados["total_frames"],
            poligono=dados["poligono"], linha=dados["linha"], classes=dados["classes"],
            config=ConfigProcessamento.from_dict(dados.get("config")),
            modelo=dados.get("modelo"), eventos=dados.get("eventos", []),
            pasta=os.path.dirname(os.path.abspath(caminho_json)),
            duracao_processamento=dados.get("duracao_processamento"),
        )
        # Se o vídeo foi movido, tenta achar ao lado do resultado
        if not os.path.isfile(res.video):
            alternativo = os.path.join(res.pasta, os.path.basename(res.video))
            if os.path.isfile(alternativo):
                res.video = alternativo
        return res

    # -------------------------------------------------------------
    # EXCEL
    # -------------------------------------------------------------

    def exportar_excel(self, caminho):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        negrito = Font(bold=True)
        cabecalho_fill = PatternFill("solid", fgColor="DDE5F0")
        duvida_fill = PatternFill("solid", fgColor="FFF2CC")
        centro = Alignment(horizontal="center")

        def cabecalho(ws, colunas):
            ws.append(colunas)
            for c in ws[ws.max_row]:
                c.font = negrito
                c.fill = cabecalho_fill
                c.alignment = centro
            ws.freeze_panes = f"A{ws.max_row + 1}"

        def larguras(ws, valores):
            for i, v in enumerate(valores, start=1):
                ws.column_dimensions[get_column_letter(i)].width = v

        # ---------- Resumo ----------
        ws = wb.active
        ws.title = "Resumo"
        ws.append(["Vídeo", os.path.basename(self.video)])
        ws.append(["Duração", formatar_tempo(self.duracao_video)])
        ws.append(["Modelo", self.modelo or ""])
        pend = len(self.duvidas(so_pendentes=True))
        ws.append(["Dúvidas não revisadas", pend])
        if pend:
            ws.cell(row=ws.max_row, column=2).fill = duvida_fill
        ws.append([])
        for linha in range(1, 5):
            ws.cell(row=linha, column=1).font = negrito

        classes = self.classes_relatorio()
        nomes = [nome_exibicao(c) for c in classes]
        cabecalho(ws, ["Início", "Fim"] + nomes + ["Total"])
        for linha in self.resumo_intervalos():
            valores = [linha["contagem"].get(c, 0) for c in classes]
            ws.append([formatar_tempo(linha["inicio"]), formatar_tempo(linha["fim"])] + valores + [sum(valores)])

        tot = self.totais()
        valores = [tot.get(c, 0) for c in classes]
        ws.append(["TOTAL", ""] + valores + [sum(valores)])
        for c in ws[ws.max_row]:
            c.font = negrito
        larguras(ws, [22, 12] + [14] * len(classes) + [10])

        # ---------- Veículos ----------
        ws = wb.create_sheet("Veículos")
        cabecalho(ws, ["Nº", "Tempo", "Classe", "Sugestão do modelo", "Confiança",
                       "Sentido", "Origem", "Dúvida", "Motivo", "Revisado", "Descartado"])
        for e in self.eventos:
            ws.append([
                e["id"], formatar_tempo(e["tempo"], com_decimos=True),
                nome_exibicao(e["classe"]),
                nome_exibicao(e["classe_sugerida"]) if e.get("classe_sugerida") else "",
                e.get("confianca"),
                {1: "1", -1: "2"}.get(e.get("sentido"), ""),
                {"automatico": "automático", "manual": "manual", "movimento": "movimento"}.get(e.get("origem"), ""),
                "sim" if e.get("duvida") else "",
                "; ".join(e.get("motivos", [])),
                "sim" if e.get("revisado") else "",
                "sim" if e.get("descartado") else "",
            ])
            if e.get("duvida") and not e.get("revisado"):
                for c in ws[ws.max_row]:
                    c.fill = duvida_fill
        larguras(ws, [7, 13, 12, 18, 11, 9, 12, 8, 45, 10, 11])

        # ---------- Por sentido ----------
        ws = wb.create_sheet("Por sentido")
        cabecalho(ws, ["Sentido"] + nomes + ["Total"])
        for sentido in (1, -1):
            cont = {c: 0 for c in classes}
            for e in self.eventos_validos:
                if e.get("sentido") == sentido:
                    cont[e["classe"]] += 1
            valores = [cont[c] for c in classes]
            ws.append([f"Sentido {1 if sentido == 1 else 2}"] + valores + [sum(valores)])
        larguras(ws, [14] + [14] * len(classes) + [10])

        wb.save(caminho)
        return caminho
